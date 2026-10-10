"""PPT 大纲导出为真实 .pptx 文件(roadmap 四十八章"后期可用 python-pptx 导出").

输入为前端持有的 PPT 大纲 JSON(与 /ai/ppt-outline/generate 输出同构),
输出标准 Office Open XML 字节流; 纯内存生成, 不落盘.
"""
from io import BytesIO

from pptx import Presentation
from pptx.util import Pt

MAX_PAGES = 30


def _set_text(placeholder, text: str, size: int | None = None) -> None:
    if placeholder is None:
        return
    placeholder.text = text
    if size is not None:
        for paragraph in placeholder.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size)


def build_pptx(outline: dict) -> bytes:
    """校验大纲结构并生成 .pptx 字节; 结构非法抛 ValueError 由路由转 40001."""
    title = str(outline.get("title") or "毕业设计答辩").strip() or "毕业设计答辩"
    pages = outline.get("pages")
    if not isinstance(pages, list) or not pages:
        raise ValueError("大纲缺少 pages, 请先生成 PPT 大纲")
    if len(pages) > MAX_PAGES:
        raise ValueError(f"大纲页数超过上限 {MAX_PAGES}")

    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    _set_text(cover.shapes.title, title, 30)
    try:
        _set_text(cover.placeholders[1], "答辩演示文稿 · AI 生成大纲", 16)
    except (KeyError, IndexError):
        pass  # 模板无副标题占位符时跳过

    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        no = page.get("page")
        heading = f"{no}. {page.get('title') or ''}" if no not in (None, "") else (page.get("title") or "")
        _set_text(slide.shapes.title, str(heading), 24)
        body = None
        try:
            body = slide.placeholders[1]
        except (KeyError, IndexError):
            pass
        if body is not None:
            points = [str(p) for p in (page.get("key_points") or []) if str(p).strip()]
            frame = body.text_frame
            frame.clear()
            if not points:
                frame.text = "(本页要点待补充)"
            else:
                frame.text = f"· {points[0]}"
                for point in points[1:]:
                    frame.add_paragraph().text = f"· {point}"
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = (
            f"页面目标: {page.get('goal') or page.get('title') or '—'}\n"
            f"讲解提示: {page.get('speaker_notes') or '—'}\n"
            f"建议图表: {page.get('recommended_visual') or '—'}"
        )

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
