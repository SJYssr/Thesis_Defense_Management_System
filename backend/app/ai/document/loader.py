"""文档解析: PDF(PyMuPDF)/DOCX/PPTX/XLSX/TXT/MD → [{section, content, page}].

文本清洗规则(roadmap 三十四章): 去页码/多余空白/无意义字符, 保留章节结构.
"""
import re
from dataclasses import dataclass, field


@dataclass
class Section:
    title: str
    content: str
    page: int | None = None
    metadata: dict = field(default_factory=dict)


def clean_text(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)  # 控制字符
    text = re.sub(r"^\s*[-—]?\s*\d+\s*[-—]?\s*$", "", text, flags=re.MULTILINE)  # 独立页码行
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


HEADING_PATTERNS = [
    re.compile(r"^(第[一二三四五六七八九十百\d]+[章节部分][^\n]{0,60})$", re.MULTILINE),
    re.compile(r"^(\d+(\.\d+){0,3})\s+[^\n]{1,80}$", re.MULTILINE),
    re.compile(r"^([一二三四五六七八九十]+、[^\n]{1,60})$", re.MULTILINE),
]


def split_by_headings(text: str) -> list[tuple[str, str]]:
    """按中文章节标题切分文本为 (标题, 正文) 列表."""
    matches: list[tuple[int, int, str]] = []
    for pattern in HEADING_PATTERNS:
        for m in pattern.finditer(text):
            matches.append((m.start(), m.end(), m.group(0).strip()))
    if not matches:
        return [("正文", text)]
    # 去重叠并按位置排序
    matches.sort(key=lambda x: x[0])
    merged: list[tuple[int, int, str]] = []
    for s, e, t in matches:
        if merged and s < merged[-1][1]:
            continue
        merged.append((s, e, t))
    sections = []
    if merged[0][0] > 0 and text[: merged[0][0]].strip():
        sections.append(("前言", text[: merged[0][0]]))
    for i, (s, e, t) in enumerate(merged):
        end = merged[i + 1][0] if i + 1 < len(merged) else len(text)
        sections.append((t, text[e:end]))
    return sections


def parse_pdf(path: str) -> list[Section]:
    import fitz

    sections: list[Section] = []
    with fitz.open(path) as doc:
        for page_no, page in enumerate(doc, start=1):
            text = clean_text(page.get_text("text"))
            if not text:
                continue
            for title, content in split_by_headings(text):
                sections.append(Section(title=title, content=content, page=page_no))
    return _merge_sections(sections)


def parse_docx(path: str) -> list[Section]:
    from docx import Document

    doc = Document(path)
    sections: list[Section] = []
    current_title, buf = "前言", []

    def flush():
        content = clean_text("\n".join(buf))
        if content:
            sections.append(Section(title=current_title, content=content))
        buf.clear()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if "heading" in style or style.startswith("标题"):
            flush()
            current_title = text
        else:
            buf.append(text)
    flush()

    for table in doc.tables:  # 表格拼接为文本
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append(Section(title="表格", content="\n".join(rows)))
    return _merge_sections(sections)


def parse_pptx(path: str) -> list[Section]:
    from pptx import Presentation

    prs = Presentation(path)
    sections = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        texts = []
        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip()
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        texts.append(t)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                texts.append(f"[备注] {note}")
        if texts:
            sections.append(Section(title=title or f"第{idx}页", content="\n".join(texts), page=idx))
    return sections


def parse_xlsx(path: str) -> list[Section]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sections = []
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows(max_row=500, values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
        if lines:
            sections.append(Section(title=f"工作表:{ws.title}", content="\n".join(lines)))
    wb.close()
    return sections


def parse_txt(path: str) -> list[Section]:
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            with open(path, encoding=enc) as f:
                text = f.read()
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:
        return []
    text = clean_text(text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)  # markdown 标题前缀
    return [Section(title=t, content=c) for t, c in split_by_headings(text)]


PARSERS = {
    ".pdf": parse_pdf, ".docx": parse_docx, ".pptx": parse_pptx,
    ".xlsx": parse_xlsx, ".txt": parse_txt, ".md": parse_txt,
}


def parse_document(path: str) -> list[Section]:
    import os

    ext = os.path.splitext(path)[1].lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"不支持的文档格式: {ext}")
    return parser(path)


def _merge_sections(sections: list[Section]) -> list[Section]:
    """合并同名相邻小节, 过滤过短片段."""
    merged: list[Section] = []
    for s in sections:
        if not s.content or len(s.content) < 10:
            continue
        if merged and merged[-1].title == s.title:
            merged[-1].content += "\n" + s.content
        else:
            merged.append(s)
    return merged
