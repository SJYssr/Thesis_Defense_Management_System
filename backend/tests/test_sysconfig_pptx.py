"""系统配置与 PPTX 导出单测(roadmap 四十八/八十五章)."""
import unittest
from io import BytesIO

from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.ai.defense.pptx_export import MAX_PAGES, build_pptx
from app.db.base import Base
from app.models import all as _all_models  # noqa: F401
from app.shared.sysconfig import SystemConfigError, _load_from_db, update_configs


def _outline(pages: int = 2) -> dict:
    return {
        "title": "测试答辩PPT",
        "pages": [
            {"page": i + 1, "title": f"第{i+1}节", "goal": f"目标{i+1}",
             "key_points": [f"要点{j}" for j in range(1 + i % 3)],
             "speaker_notes": "讲稿", "recommended_visual": "架构图"}
            for i in range(pages)
        ],
    }


class PptxExportTest(unittest.TestCase):
    def test_build_valid_outline(self):
        content = build_pptx(_outline(3))
        self.assertTrue(content[:2] == b"PK", "应为 zip(OOXML) 头")
        prs = Presentation(BytesIO(content))
        self.assertEqual(len(prs.slides), 4, "封面+3 内容页")
        notes = prs.slides[1].notes_slide.notes_text_frame.text
        self.assertIn("页面目标: 目标1", notes)
        self.assertIn("建议图表: 架构图", notes)

    def test_empty_and_oversize_rejected(self):
        with self.assertRaises(ValueError):
            build_pptx({"title": "空", "pages": []})
        with self.assertRaises(ValueError):
            build_pptx(_outline(MAX_PAGES + 1))

    def test_missing_optional_fields_tolerated(self):
        content = build_pptx({"pages": [{"title": "仅标题"}]})
        prs = Presentation(BytesIO(content))
        self.assertEqual(len(prs.slides), 2)


class SysConfigTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_defaults_without_rows(self):
        values = _load_from_db(self.db)
        self.assertEqual(values["system_name"], "毕业设计管理系统")
        self.assertEqual(values["ai_defense_enabled"], "true")

    def test_update_and_persist(self):
        values = update_configs(self.db, {"system_name": "新名称", "ai_defense_enabled": "false"}, updater_id=1)
        self.assertEqual(values["system_name"], "新名称")
        # 独立会话读取验证已持久化
        with Session(self.engine) as other:
            self.assertEqual(_load_from_db(other)["ai_defense_enabled"], "false")

    def test_validation(self):
        with self.assertRaises(SystemConfigError):
            update_configs(self.db, {"unknown_key": "x"}, updater_id=1)
        with self.assertRaises(SystemConfigError):
            update_configs(self.db, {"ai_defense_enabled": "yes"}, updater_id=1)
        with self.assertRaises(SystemConfigError):
            update_configs(self.db, {"ai_rate_limit_per_minute": "0"}, updater_id=1)
        with self.assertRaises(SystemConfigError):
            update_configs(self.db, {"max_upload_size_mb": "abc"}, updater_id=1)


if __name__ == "__main__":
    unittest.main()
